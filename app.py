import io
import re
import os
import zipfile
import datetime
import base64
import fitz  # PyMuPDF
import pdfplumber
import pandas as pd
import numpy as np
import streamlit as st
from PIL import Image
from io import BytesIO
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ==========================================
# CARGA DINÁMICA DEL LOGO (FAVICON) - OPCIÓN RECOMENDADA
# ==========================================
icono_pagina = "🏢"
archivos_logo_posibles = [
    "LOGO SANTANDER.jpg",
    "LOGO ZFS-ZFC.jpeg",
    "logo.jpeg",
    "logo.jpg",
    "logo.png",
    "favicon.ico"
]

for archivo in archivos_logo_posibles:
    if os.path.exists(archivo):
        try:
            icono_pagina = Image.open(archivo)
            break
        except Exception:
            pass

# ==========================================
# CONFIGURACIÓN GENERAL Y ESTILOS CORPORATIVOS
# ==========================================
st.set_page_config(
    page_title="Sistema de Automatización Aduanera | Zona Franca",
    page_icon=icono_pagina,
    layout="wide",
    initial_sidebar_state="expanded"
)

fondo_css = ""
if os.path.exists("Fondo ZFC.png"):
    with open("Fondo ZFC.png", "rb") as f:
        fondo_bytes = f.read()
    fondo_base64 = base64.b64encode(fondo_bytes).decode()
    fondo_css = f"""
    .stApp {{
        background-image: linear-gradient(rgba(244, 247, 246, 0.9), rgba(244, 247, 246, 0.9)), url("data:image/png;base64,{fondo_base64}");
        background-size: cover;
        background-position: center;
        background-repeat: no-repeat;
    }}
    """
else:
    fondo_css = ".stApp { background-color: #F8F9FA; }"

st.markdown(f"""
    <style>
    :root {{
        --zf-green-dark: #12402A;
        --zf-green-medium: #1F4E3D;
        --zf-olive: #8A9A28;
        --zf-card-bg: #FFFFFF;
        --zf-text-main: #2C3E50;
    }}
    {fondo_css}

    .executive-title {{
        font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
        color: var(--zf-green-dark);
        font-size: 26px;
        font-weight: 700;
        margin-bottom: 0px;
        letter-spacing: -0.5px;
    }}
    .executive-subtitle {{
        font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
        color: #4B5563;
        font-size: 14px;
        margin-top: 5px;
        margin-bottom: 20px;
    }}
    [data-testid="stSidebar"] {{
        background-color: #F3F4F6;
        border-right: 1px solid #E5E7EB;
    }}
    .stButton>button {{
        background-color: var(--zf-green-dark);
        color: white;
        border-radius: 6px;
        border: none;
        font-weight: 600;
        padding: 0.5rem 1rem;
        box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        transition: all 0.3s ease;
    }}
    .stButton>button:hover {{
        background-color: var(--zf-green-medium);
        color: white;
        border: none;
    }}
    </style>
""", unsafe_allow_html=True)

# ESTILOS OPENPYXL COMPARTIDOS
HEADER_FILL = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill(start_color="D9E1D9", end_color="D9E1D9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

# --- NAVEGACIÓN EN LA BARRA LATERAL ---
with st.sidebar:
    logo_encontrado = False
    for filename in archivos_logo_posibles:
        if os.path.exists(filename) and filename != "favicon.ico":
            st.image(filename, use_container_width=True)
            logo_encontrado = True
            break
    if not logo_encontrado:
        st.markdown("<h3 style='color: #12402A; text-align: center;'>ZONA FRANCA</h3>", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("**Navegación Operativa**")
    st.caption("Seleccione el módulo de gestión:")

    opcion_modulo = st.radio(
        label="Módulo de Gestión",
        options=[
            "Procesador de DIM",
            "Actas de Tránsito / PICIZ",
            "Control Bloqueo Formularios"
        ],
        label_visibility="collapsed"
    )

    st.markdown("<hr style='margin: 20px 0 10px 0; border: 1px solid #E5E7EB;'>", unsafe_allow_html=True)
    st.markdown("""
        <div style='font-size: 12px; color: #6B7280; line-height: 1.4;'>
            <b>Zona Franca de Cúcuta</b><br>
            <i>Operada por Zona Franca Santander</i><br>
            Sistema de Automatización Aduanera
        </div>
    """, unsafe_allow_html=True)


# ==========================================
# MÓDULO 1: PROCESADOR DE DIM (FORMULARIO 500)
# ==========================================
COLUMNAS_DIM = [
    "Número de formulario", "NIT Importador", "Razón Social Importador", "Factura",
    "Manifiesto de carga", "Documento de transporte", "Cod. País Procedencia",
    "Cod. Modo Transporte", "Código de Bandera", "Tasa de Cambio", "Subpartida Arancelaria",
    "Cod. País Origen", "Cod. País Compra", "Peso Bruto (Kgs)", "Peso Neto (Kgs)",
    "Código de Embalaje", "Cod. Unidad Comercial (76)", "Cantidad (77)", "No. Bultos",
    "Valor FOB (USD)", "Sumatoria Fletes/Seguros/Otros (USD)", "Acta de Inspección No.",
    "Levante No.", "Fecha del Levante", "Archivo", "Campos_no_encontrados"
]
MONTO = r"[\d\.,]+"
ENTERO_MILES = r"[\d\.,]+"

def limpiar_monto_dim(val_str):
    if not val_str:
        return 0.0
    val_str = str(val_str).strip()
    val_str = re.sub(r'[^\d.,]', '', val_str)
    if not val_str:
        return 0.0

    if ',' in val_str:
        val_str = val_str.replace('.', '').replace(',', '.')
    else:
        partes = val_str.split('.')
        if len(partes) == 2 and len(partes[1]) == 3 and len(partes[0]) <= 3:
            val_str = val_str.replace('.', '')
        elif len(partes) > 2:
            val_str = "".join(partes[:-1]) + "." + partes[-1] if len(partes[-1]) == 2 else val_str.replace('.', '')

    try:
        return float(val_str)
    except ValueError:
        return 0.0

def _buscar_dim(patron, texto, flags=re.IGNORECASE, grupo=1):
    m = re.search(patron, texto, flags)
    return m.group(grupo).strip() if m else None

def dividir_dims(texto: str):
    partes = re.split(r"(?=Declaraci[oó]n de Importaci[oó]n)", texto, flags=re.IGNORECASE)
    return [p for p in partes if re.search(r"N[uú]mero de formulario", p, re.IGNORECASE)]

def extraer_campos_dim(chunk_texto: str, texto_completo: str, nombre_archivo: str) -> dict:
    faltantes = []
    def campo(nombre, patron, grupo=1, default=""):
        valor = _buscar_dim(patron, chunk_texto, grupo=grupo) or _buscar_dim(patron, texto_completo, grupo=grupo)
        if not valor:
            faltantes.append(nombre)
            return default
        return " ".join(valor.split())

    numero_formulario = campo("Número de formulario", r"4\s*\.\s*N[uú]mero de formulario\s*\n?\s*(\S+)")
    nit_importador = campo("NIT Importador", r"5\s*\.\s*N[uú]mero de Identificaci[oó]n Tributaria \(NIT\)\s*(\d{9,10})")
    razon_social = campo("Razón Social Importador", r"11\s*\.\s*Apellidos y nombres o Raz[oó]n Social\s*([^\n]+)")
    factura = campo("Factura", r"51\s*\.\s*No\.\s*de\s*factura\s*\n\s*(\S+)")
    
    manifiesto_carga = campo("Manifiesto de carga", r"42\s*\.?\s*Manifiesto\s+de\s+carga\s*(?:No\.?\s*)?([A-Za-z0-9\-]+)") or campo("Manifiesto de carga", r"42\s*\.?\s*Manifiesto\s+de\s+carga[^\n]*\n\s*(?:No\.?\s*)?([A-Za-z0-9\-]+)")
    documento_transporte = campo("Documento de transporte", r"44\s*\.?\s*Documento\s+de\s+transporte\s*(?:No\.?\s*)?([A-Za-z0-9\-]+)") or campo("Documento de transporte", r"44\s*\.?\s*Documento\s+de\s+transporte[^\n]*\n\s*(?:No\.?\s*)?([A-Za-z0-9\-]+)")

    cod_pais_procedencia = campo("Cod. País Procedencia", r"53\s*\.\s*(?:C[oó]d\.?\s*)?pa[ií]s\s+(?:de\s+)?procedencia\s*([A-Za-z0-9]{2,3})")
    cod_modo_transporte = campo("Cod. Modo Transporte", r"54\s*\.\s*Cod\.\s*Modo\s*Transporte\s*(\d)")
    codigo_bandera = campo("Código de Bandera", r"55\s*\.\s*C[oó]digo\s+(?:de\s+)?bandera\s*([A-Za-z0-9]{2,3})")
    tasa_cambio = campo("Tasa de Cambio", r"Tasa de cambio\s*\$?\s*cvs\.?\s*\n?\s*([\d.,]+)")
    subpartida = campo("Subpartida Arancelaria", r"59\s*\.\s*Subpartida arancelaria\s*(\d{10})")
    cod_pais_origen = campo("Cod. País Origen", r"66\s*\.\s*(?:C[oó]d\.?\s*)?pa[ií]s\s+(?:de\s+)?origen\s*([A-Za-z0-9]{2,3})")
    cod_pais_compra = campo("Cod. País Compra", r"70\s*\.\s*Cod\s*\.\s*pa[ií]s\s*\n?\s*compra\s*(\d{2,3})")
    codigo_embalaje = campo("Código de Embalaje", r"73\s*\.\s*C[oó]digo\s*\n?\s*embalaje\s*([A-Za-z0-9]{1,4})")
    
    peso_bruto = limpiar_monto_dim(campo("Peso Bruto (Kgs)", r"71\s*\.\s*Peso bruto kgs\.\s*dcms\.\s*(" + MONTO + ")"))
    peso_neto = limpiar_monto_dim(campo("Peso Neto (Kgs)", r"72\s*\.\s*Peso neto kgs\.\s*dcms\.\s*(" + MONTO + ")"))
    valor_fob = limpiar_monto_dim(campo("Valor FOB (USD)", r"78\s*\.\s*Valor FOB USD\s*(" + MONTO + ")"))
    sumatoria_fletes = limpiar_monto_dim(campo("Sumatoria Fletes/Seguros/Otros (USD)", r"82\s*\.\s*Sumatoria de fletes,?\s*seguros\s*\n?\s*y otros gastos USD\s*(" + MONTO + ")"))
    
    cod_unidad_comercial = campo("Cod. Unidad Comercial (76)", r"76\s*\.\s*Cod\.?\s*unidad\s*comercial\s*[\r\n\s]+([A-Za-z]{1,4})\b") or campo("Cod. Unidad Comercial (76)", r"76\s*\.\s*Cod\.?\s*unidad\s*comercial\s+([A-Za-z]{1,4})\b")
    cantidad_str = campo("Cantidad (77)", r"77\s*\.\s*Cantidad\s*(?:dcms\.?)?\s*[\r\n\s]+(" + MONTO + ")") or campo("Cantidad (77)", r"77\s*\.\s*Cantidad\s*(?:dcms\.?)?\s*(" + MONTO + ")")
    cantidad_comercial = limpiar_monto_dim(cantidad_str)

    n_bultos_str = campo("No. Bultos", r"74\s*\.\s*No\.\s*bultos\s*(" + ENTERO_MILES + ")")
    try:
        no_bultos = int(re.sub(r'[^\d]', '', n_bultos_str)) if n_bultos_str else 0
    except ValueError:
        no_bultos = 0

    m_acta = re.search(r"ACTA\s+DE\s+INSPECCI[OÓ]N\s*(?:No\.?|Número)?\s*[:\.]?\s*([0-9]{8,15})", texto_completo, re.IGNORECASE)
    acta_inspeccion = m_acta.group(1).strip() if m_acta else ""

    m_lev_box = re.search(r"134\.?\s*Levante\s+No\.?\s*([0-9]{8,15})", chunk_texto, re.IGNORECASE) or re.search(r"(?:Levante|Auto(?:rización)?)\s*(?:No\.?|Número)?\s*[:\.]?\s*([0-9]{8,15})", texto_completo, re.IGNORECASE)
    levante_no = m_lev_box.group(1).strip() if m_lev_box else ""
    if not levante_no:
        faltantes.append("Levante No.")

    fecha_levante = campo("Fecha del Levante", r"135\.?\s*Fecha[^\d\n]*(\d{4}\s*[-/\.]\s*\d{2}\s*[-/\.]\s*\d{2})")
    if not fecha_levante:
        m_fec = re.search(r"\b(20\d{2}[-/\.](?:0[1-9]|1[0-2])[-/\.](?:0[1-9]|[12]\d|3[01]))\b", texto_completo)
        if m_fec:
            fecha_levante = m_fec.group(1)
            if "Fecha del Levante" in faltantes:
                faltantes.remove("Fecha del Levante")

    if fecha_levante:
        fecha_levante = re.sub(r"\s+", "", fecha_levante)
        fecha_levante = re.sub(r"[/.]", "-", fecha_levante)

    return {
        "Número de formulario": numero_formulario, "NIT Importador": nit_importador,
        "Razón Social Importador": razon_social, "Factura": factura,
        "Manifiesto de carga": manifiesto_carga, "Documento de transporte": documento_transporte,
        "Cod. País Procedencia": cod_pais_procedencia, "Cod. Modo Transporte": cod_modo_transporte,
        "Código de Bandera": codigo_bandera, "Tasa de Cambio": tasa_cambio,
        "Subpartida Arancelaria": subpartida, "Cod. País Origen": cod_pais_origen,
        "Cod. País Compra": cod_pais_compra, "Peso Bruto (Kgs)": peso_bruto,
        "Peso Neto (Kgs)": peso_neto, "Código de Embalaje": codigo_embalaje,
        "Cod. Unidad Comercial (76)": cod_unidad_comercial, "Cantidad (77)": cantidad_comercial,
        "No. Bultos": no_bultos, "Valor FOB (USD)": valor_fob,
        "Sumatoria Fletes/Seguros/Otros (USD)": sumatoria_fletes,
        "Acta de Inspección No.": acta_inspeccion, "Levante No.": levante_no,
        "Fecha del Levante": fecha_levante, "Archivo": nombre_archivo,
        "Campos_no_encontrados": ", ".join(faltantes) if faltantes else ""
    }

def extraer_texto_pdf(data: bytes) -> str:
    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)

def obtener_pdfs_desde_upload(uploaded_file):
    nombre, contenido = uploaded_file.name, uploaded_file.read()
    if nombre.lower().endswith(".zip"):
        pdfs = []
        with zipfile.ZipFile(io.BytesIO(contenido)) as zf:
            for info in zf.infolist():
                if info.filename.lower().endswith(".pdf") and not info.is_dir():
                    pdfs.append((info.filename.split("/")[-1], zf.read(info.filename)))
        return pdfs
    elif nombre.lower().endswith(".pdf"):
        return [(nombre, contenido)]
    return []

def procesar_archivos_dim(uploaded_files, progress_callback=None) -> pd.DataFrame:
    filas, tareas = [], []
    for uf in uploaded_files:
        tareas.extend(obtener_pdfs_desde_upload(uf))

    total = max(len(tareas), 1)
    for i, (nombre_pdf, data) in enumerate(tareas, start=1):
        try:
            texto_completo = extraer_texto_pdf(data)
            dim_chunks = dividir_dims(texto_completo) or [texto_completo]
            for chunk in dim_chunks:
                filas.append(extraer_campos_dim(chunk, texto_completo, nombre_pdf))
        except Exception as exc: 
            fila = {c: "" for c in COLUMNAS_DIM}
            fila["Archivo"], fila["Campos_no_encontrados"] = nombre_pdf, f"ERROR: {exc}"
            filas.append(fila)
        if progress_callback:
            progress_callback(i / total, nombre_pdf)

    if not filas:
        return pd.DataFrame(columns=COLUMNAS_DIM)

    df = pd.DataFrame(filas, columns=COLUMNAS_DIM)
    if "Número de formulario" in df.columns:
        df["Número de formulario"] = df["Número de formulario"].astype(str).str.strip()
        df = df[df["Número de formulario"].notna() & (df["Número de formulario"] != "") & (df["Número de formulario"].str.lower() != "nan")]
        df = df.drop_duplicates(subset=["Número de formulario"], keep="last").reset_index(drop=True)
    return df

def generar_excel_dim(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    df_export = df.copy()
    if not df_export.empty:
        total_row = {c: "" for c in df_export.columns}
        total_row["Número de formulario"] = "TOTALES CONSOLIDADOS"
        total_row["Valor FOB (USD)"] = df_export["Valor FOB (USD)"].sum()
        total_row["Sumatoria Fletes/Seguros/Otros (USD)"] = df_export["Sumatoria Fletes/Seguros/Otros (USD)"].sum()
        total_row["Peso Bruto (Kgs)"] = df_export["Peso Bruto (Kgs)"].sum()
        total_row["Peso Neto (Kgs)"] = df_export["Peso Neto (Kgs)"].sum()
        total_row["No. Bultos"] = df_export["No. Bultos"].sum()
        if "Cantidad (77)" in df_export.columns:
            total_row["Cantidad (77)"] = df_export["Cantidad (77)"].sum()
        df_export = pd.concat([df_export, pd.DataFrame([total_row])], ignore_index=True)

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="DIM")
        ws = writer.sheets["DIM"]
        n_filas, n_cols = df_export.shape[0], df_export.shape[1]

        for col_idx in range(1, n_cols + 1):
            celda = ws.cell(row=1, column=col_idx)
            celda.fill, celda.font, celda.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx in range(2, n_filas + 2):
            for col_idx in range(1, n_cols + 1):
                celda = ws.cell(row=row_idx, column=col_idx)
                celda.border = THIN_BORDER
                celda.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(celda.value, (int, float)):
                    celda.number_format = '#,##0.00'
                if row_idx == n_filas + 1:
                    celda.fill, celda.font = TOTAL_FILL, Font(bold=True)

        for col_idx, columna in enumerate(df_export.columns, start=1):
            longitudes = [len(str(columna))] + [len(str(v)) for v in df_export[columna].astype(str).tolist()]
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(longitudes) + 3, 45)
        ws.freeze_panes = "A2"
    return buffer.getvalue()

def modulo_procesador_dim():
    st.markdown("<p class='executive-title'>Procesador de Declaración de Importación (DIM)</p>", unsafe_allow_html=True)
    st.markdown("<p class='executive-subtitle'>Módulo masivo para extracción y consolidación de datos del Formulario 500 DIAN.</p>", unsafe_allow_html=True)
    
    if "df_resultado_dim" not in st.session_state:
        st.session_state.df_resultado_dim = pd.DataFrame(columns=COLUMNAS_DIM)
    if "uploader_key_dim" not in st.session_state:
        st.session_state.uploader_key_dim = 0

    uploaded_files = st.file_uploader(
        "Carga tus archivos PDF individuales o paquetes comprimidos .ZIP",
        type=["pdf", "zip"],
        accept_multiple_files=True,
        key=f"dim_uploader_{st.session_state.uploader_key_dim}"
    )

    col_btn1, col_btn2 = st.columns(2)
    with col_btn1:
        procesar = st.button("🚀 Procesar Documentos DIM", type="primary", use_container_width=True)
    with col_btn2:
        limpiar = st.button("🧹 Limpiar Panel DIM", use_container_width=True)

    if limpiar:
        st.session_state.df_resultado_dim = pd.DataFrame(columns=COLUMNAS_DIM)
        st.session_state.uploader_key_dim += 1
        st.rerun()

    if procesar:
        if not uploaded_files:
            st.warning("Por favor carga al menos un archivo PDF o ZIP.")
        else:
            progreso = st.progress(0.0, text="Iniciando motor de extracción...")
            df = procesar_archivos_dim(uploaded_files, progress_callback=lambda pct, nom: progreso.progress(pct, text=f"Procesando: {nom}"))
            progreso.empty()
            st.session_state.df_resultado_dim = df
            st.success(f"✅ Extracción completada. {len(df)} declaración(es) procesada(s).")

    df = st.session_state.df_resultado_dim
    if not df.empty:
        st.markdown("---")
        st.subheader("Consolidado y Analítica de Carga")
        cols_totales = st.columns(4)
        cols_totales[0].metric("Total Valor FOB (USD)", f"${df['Valor FOB (USD)'].sum():,.2f}")
        cols_totales[1].metric("Total Fletes/Seguros (USD)", f"${df['Sumatoria Fletes/Seguros/Otros (USD)'].sum():,.2f}")
        cols_totales[2].metric("Total Peso Bruto (Kgs)", f"{df['Peso Bruto (Kgs)'].sum():,.2f}")
        cols_totales[3].metric("Total Peso Neto (Kgs)", f"{df['Peso Neto (Kgs)'].sum():,.2f}")

        st.markdown("<br>", unsafe_allow_html=True)
        busqueda = st.text_input("🔍 Búsqueda rápida en el consolidado...", "")
        df_vista = df.copy()
        if busqueda:
            mask = df_vista.apply(lambda fila: fila.astype(str).str.contains(busqueda, case=False, na=False).any(), axis=1)
            df_vista = df_vista[mask]

        st.dataframe(df_vista, use_container_width=True, height=400)
        
        st.markdown("---")
        col1, col2 = st.columns(2)
        with col1:
            st.download_button(
                "⬇️ Descargar Reporte Excel (.xlsx)",
                data=generar_excel_dim(df.drop(columns=["Campos_no_encontrados"])),
                file_name=f"dim_zona_franca_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with col2:
            st.download_button(
                "⬇️ Descargar Reporte CSV (.csv)",
                data=df.drop(columns=["Campos_no_encontrados"]).to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig"),
                file_name=f"dim_zona_franca_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                mime="text/csv",
                use_container_width=True
            )


# ==========================================
# MÓDULO 2: ACTAS DE TRÁNSITO / PICIZ
# ==========================================
def extraer_datos_acta(pdf_file, nombre_archivo):
    with pdfplumber.open(pdf_file) as pdf:
        texto = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])

    usuario = re.search(r"consignados?\s+al\s*\n?\s*([A-Z0-9\.\-\s]+?)(?=\s+y\s+amparado|\s+DECLARACION|\n\s*DECLARACION|$)", texto, re.IGNORECASE)
    val_usuario = " ".join(usuario.group(1).strip().split()) if usuario else "N/A"

    doc_form = re.search(r"DOCUMENTO\s+FORMULARIO\s+MERCANC[ÍI]A[\s\S]*?\n\s*([A-Z0-9\.\-_]+)\s+(\d+)", texto, re.IGNORECASE) or re.search(r"\b([A-Z0-9\.\-_]{5,})\s+(\d{7,10})\b", texto)
    transito = re.search(r"DECLARACION\s+DE\s+TRANSITO\s+ADUANERO\s*\n?\s*(?:Número|N[úu]mero)?\s*[:\.]?\s*(\d+)", texto, re.IGNORECASE)
    fecha_ingreso = re.search(r"ACTA\s+DE\s+DESPRECINTAJE[\s\S]*?(\d{2}/\d{2}/\d{4})", texto, re.IGNORECASE)
    fecha_auto = re.search(r"Fecha\s+de\s+la\s+autorizaci[oó]n\s+de\s+la\s+operaci[oó]n[^\d]*(\d{4}[-/]\d{2}[-/]\d{2})", texto, re.IGNORECASE)
    fecha_limite = re.search(r"Fecha\s+l[ií]mite\s+para\s+finalizar\s+el\s+r[eé]gimen[^\d]*(\d{4}[-/]\d{2}[-/]\d{2})", texto, re.IGNORECASE)
    acta_n = re.search(r"Acta\s+N\.?\s*(\d+)", texto, re.IGNORECASE)
    fecha_acta_m = re.search(r"FECHA\s+GENERACI[OÓ]N\s+DEL\s+ACTA:\s*(\d{2}/\d{2}/\d{4})", texto, re.IGNORECASE)
    fecha_acta = fecha_acta_m.group(1).strip() if fecha_acta_m else "N/A"

    peso_match = re.search(r"TOTALES\s*:\s*[\d\.,]+\s+([\d\.,]+)", texto, re.IGNORECASE) or re.search(r"TOTALES\s*:\s*([\d\.,]+)", texto, re.IGNORECASE)
    peso_bascula = peso_match.group(1).strip() if peso_match else "N/A"

    obs_match = re.search(r"Observaciones[\s\S]*?\n([\s\S]*?)(?=\n\s*(?:DOCUMENTO|TOTALES|USUARIO\s+OPERADOR|FECHA\s+GENERACI|\Z))", texto, re.IGNORECASE)
    if obs_match:
        lineas_limpias = [l.strip() for l in obs_match.group(1).split("\n") if l.strip() and not re.match(r"^(Descripción\s*N/A|Bultos|Estado|Términos|Otra)\b", l.strip(), re.IGNORECASE)]
        observaciones = " ".join(lineas_limpias) if lineas_limpias else "N/A"
    else:
        observaciones = "N/A"

    return {
        "Usuario": val_usuario,
        "Documento de transporte": doc_form.group(1).strip() if doc_form else "N/A",
        "Transito N°": transito.group(1).strip() if transito else "N/A",
        "FMM N°": doc_form.group(2).strip() if doc_form else "N/A",
        "FECHA INGRESO ÚLTIMO VEHÍCULO": fecha_ingreso.group(1).strip() if fecha_ingreso else "N/A",
        "Fecha de autorización": fecha_auto.group(1).strip() if fecha_auto else "N/A",
        "Tránsito Fecha Maxima Finalización": fecha_limite.group(1).strip() if fecha_limite else "N/A",
        "Acta de Inventario e Inconsistencias PICIZ": acta_n.group(1).strip() if acta_n else "N/A",
        "Fecha acta de inventario e inconsistencias": fecha_acta,
        "No. Planilla de Recepción (FECHA)": fecha_acta,
        "Peso Báscula ZFC": peso_bascula,
        "OBSERVACIONES/ INCONSISTENCIAS": observaciones,
        "Archivo": nombre_archivo
    }

def generar_excel_actas(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Actas")
        ws = writer.sheets["Actas"]
        n_filas, n_cols = df.shape[0], df.shape[1]

        for col_idx in range(1, n_cols + 1):
            celda = ws.cell(row=1, column=col_idx)
            celda.fill, celda.font, celda.border = HEADER_FILL, HEADER_FONT, THIN_BORDER
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx in range(2, n_filas + 2):
            for col_idx in range(1, n_cols + 1):
                celda = ws.cell(row=row_idx, column=col_idx)
                celda.border = THIN_BORDER
                celda.alignment = Alignment(vertical="top", wrap_text=True)

        for col_idx, columna in enumerate(df.columns, start=1):
            longitudes = [len(str(columna))] + [len(str(v)) for v in df[columna].astype(str).tolist()]
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(longitudes) + 3, 45)
        ws.freeze_panes = "A2"
    return buffer.getvalue()

def modulo_actas_transito():
    st.markdown("<p class='executive-title'>Actas de Tránsito / PICIZ</p>", unsafe_allow_html=True)
    st.markdown("<p class='executive-subtitle'>Extractor automático de Actas de Inventario e Inconsistencias para Tránsito Aduanero.</p>", unsafe_allow_html=True)

    if "df_resultado_actas" not in st.session_state:
        st.session_state.df_resultado_actas = pd.DataFrame()
    if "uploader_key_actas" not in st.session_state:
        st.session_state.uploader_key_actas = 0

    uploaded_files = st.file_uploader(
        "Carga los archivos PDF de las actas de inventario",
        type=["pdf"],
        accept_multiple_files=True,
        key=f"actas_uploader_{st.session_state.uploader_key_actas}"
    )

    col_btn1, col_btn2 = st.columns(2)
    with col_btn1:
        procesar = st.button("🚀 Procesar Actas", type="primary", use_container_width=True)
    with col_btn2:
        limpiar = st.button("🧹 Limpiar Panel Actas", use_container_width=True)

    if limpiar:
        st.session_state.df_resultado_actas = pd.DataFrame()
        st.session_state.uploader_key_actas += 1
        st.rerun()

    if procesar:
        if not uploaded_files:
            st.warning("Por favor carga al menos un archivo PDF.")
        else:
            with st.spinner("Extrayendo datos de las actas..."):
                datos = [extraer_datos_acta(file, file.name) for file in uploaded_files]
                st.session_state.df_resultado_actas = pd.DataFrame(datos)
            st.success(f"✅ Extracción completada para {len(uploaded_files)} acta(s).")

    df = st.session_state.df_resultado_actas
    if not df.empty:
        st.markdown("---")
        busqueda = st.text_input("🔍 Búsqueda rápida en actas...", "")
        df_vista = df.copy()
        if busqueda:
            mask = df_vista.apply(lambda fila: fila.astype(str).str.contains(busqueda, case=False, na=False).any(), axis=1)
            df_vista = df_vista[mask]

        st.dataframe(df_vista, use_container_width=True, height=400)

        st.markdown("---")
        col1, col2 = st.columns(2)
        with col1:
            st.download_button(
                "⬇️ Descargar Reporte Excel (.xlsx)",
                data=generar_excel_actas(df),
                file_name=f"actas_zona_franca_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with col2:
            st.download_button(
                "⬇️ Descargar Reporte CSV (.csv)",
                data=df.to_csv(index=False, encoding="utf-8-sig").encode("utf-8-sig"),
                file_name=f"actas_zona_franca_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                mime="text/csv",
                use_container_width=True
            )


# ==========================================
# MÓDULO 3: CONTROL BLOQUEO FORMULARIOS
# ==========================================
def limpiar_numeros(valor):
    if pd.isna(valor) or valor == '':
        return ""
    try:
        val_str = str(valor).strip()
        if val_str == '' or val_str.lower() in ['nan', 'none']:
            return ""
        if '.' in val_str:
            return str(int(float(val_str)))
        return val_str
    except ValueError:
        return str(valor).strip()

def modulo_control_bloqueos():
    st.markdown("<p class='executive-title'>Módulo de Control de Bloqueos y Vencimientos</p>", unsafe_allow_html=True)
    st.markdown("<p class='executive-subtitle'>Sistema de auditoría y seguimiento de ingresos para operaciones de comercio exterior.</p>", unsafe_allow_html=True)

    if 'file_processed' not in st.session_state:
        st.session_state.file_processed = False
    if 'df_final' not in st.session_state:
        st.session_state.df_final = None
    if 'fecha_ref' not in st.session_state:
        st.session_state.fecha_ref = datetime.date.today()

    with st.sidebar:
        st.markdown("<hr style='margin: 10px 0; border: 1px solid #D1D5DB;'>", unsafe_allow_html=True)
        st.markdown("### Configuración Operativa")
        
        fecha_actual_input = st.date_input(
            "Fecha Actual de Referencia",
            value=st.session_state.fecha_ref
        )
        st.session_state.fecha_ref = fecha_actual_input

        st.markdown("---")
        uploaded_file = st.file_uploader("Cargar Reporte PW (Excel)", type=["xlsx", "xls"], key="uploader_bloqueos")
        
        st.markdown("---")
        if st.button("Limpiar Datos y Sesión", use_container_width=True, key="btn_clean_bloqueos"):
            st.session_state.file_processed = False
            st.session_state.df_final = None
            st.rerun()

    if uploaded_file is not None:
        try:
            with st.spinner("Procesando registros aduaneros..."):
                df_raw = pd.read_excel(uploaded_file, header=None, dtype=str)
                
                header_row_index = -1
                for i, row in df_raw.iterrows():
                    row_str = " ".join(str(val) for val in row.values)
                    if "NOMBRE COMPANIA" in row_str or "PLACA" in row_str:
                        header_row_index = i
                        break
                
                if header_row_index == -1:
                    st.error("No se localizó la fila de encabezados estándar en el archivo cargado.")
                    st.stop()
                
                df = pd.read_excel(uploaded_file, header=header_row_index, dtype=str)
                df.columns = df.columns.str.strip().str.replace(r'\r\n', '', regex=True)
                
                columnas_esperadas = {
                    'NOMBRE COMPANIA': 'Compañía Usuaria',
                    'PLACA': 'Placa',
                    'FECHA REGISTRO': 'Fecha Registro',
                    'TIPO INGRESO': 'Tipo Ingreso', 
                    'NUM DEL DOC. ADUANERO': 'Número Documento',
                    'TRANSITO': 'Tránsito',
                    'FECHA BASCULA': 'Fecha de Báscula'
                }
                
                columnas_existentes = {k: v for k, v in columnas_esperadas.items() if k in df.columns}
                df_filtrado = df[list(columnas_existentes.keys())].rename(columns=columnas_existentes)
                
                if 'Tipo Ingreso' in df_filtrado.columns:
                    df_filtrado = df_filtrado[df_filtrado['Tipo Ingreso'].str.contains('FORMULARIO', na=False, case=False)]
                    df_filtrado = df_filtrado.drop(columns=['Tipo Ingreso'])
                
                df_filtrado = df_filtrado.dropna(subset=['Placa'])
                df_filtrado['Placa'] = df_filtrado['Placa'].astype(str).str.strip()
                df_filtrado = df_filtrado[df_filtrado['Placa'] != '']

                for col in ['Número Documento', 'Tránsito']:
                    if col in df_filtrado.columns:
                        df_filtrado[col] = df_filtrado[col].apply(limpiar_numeros)

                for col in ['Fecha Registro', 'Fecha de Báscula']:
                    if col in df_filtrado.columns:
                        df_filtrado[col] = pd.to_datetime(df_filtrado[col], dayfirst=True, errors='coerce').dt.date

                columnas_dedup = [col for col in ['Placa', 'Fecha Registro', 'Compañía Usuaria', 'Número Documento', 'Tránsito', 'Fecha de Báscula'] if col in df_filtrado.columns]
                df_filtrado = df_filtrado.drop_duplicates(subset=columnas_dedup, keep='first')

                df_filtrado['Límite'] = 5
                
                def calcular_vencimiento(row):
                    fecha_base = row['Fecha de Báscula'] if pd.notna(row['Fecha de Báscula']) and str(row['Fecha de Báscula']) != 'NaT' else row['Fecha Registro']
                    if pd.isna(fecha_base) or str(fecha_base) == 'NaT':
                        return None
                    try:
                        fecha_venc = np.busday_offset(np.datetime64(fecha_base), 5, roll='forward')
                        return pd.to_datetime(fecha_venc).date()
                    except:
                        return None
                        
                df_filtrado['Vencimiento (5 Días Hábiles)'] = df_filtrado.apply(calcular_vencimiento, axis=1)
                
                def calcular_dias_restantes(fecha_venc):
                    if pd.isna(fecha_venc) or str(fecha_venc) == 'NaT':
                        return None
                    delta = fecha_venc - st.session_state.fecha_ref
                    return delta.days
                    
                df_filtrado['Días Restantes'] = df_filtrado['Vencimiento (5 Días Hábiles)'].apply(calcular_dias_restantes)
                
                orden_columnas = ['Placa', 'Fecha Registro', 'Compañía Usuaria', 'Número Documento', 'Tránsito', 'Fecha de Báscula', 'Límite', 'Vencimiento (5 Días Hábiles)', 'Días Restantes']
                orden_columnas = [col for col in orden_columnas if col in df_filtrado.columns]
                df_final = df_filtrado[orden_columnas].copy()
                
                df_final['Días Restantes'] = pd.to_numeric(df_final['Días Restantes'], errors='coerce').fillna(0).astype(int)
                st.session_state.df_final = df_final
                st.session_state.file_processed = True

        except Exception as e:
            st.error(f"Error en el procesamiento del archivo: {e}")

    if st.session_state.file_processed and st.session_state.df_final is not None:
        df_res = st.session_state.df_final
        
        dias_series = df_res['Días Restantes']
        vencidos = (dias_series <= 0).sum()
        riesgo = ((dias_series >= 1) & (dias_series <= 2)).sum()
        a_tiempo = (dias_series >= 3).sum()
        
        st.markdown(f"<p style='font-size: 15px; font-weight: 600; color: #12402A;'>Resumen de Estado Operativo — Fecha de Corte: {st.session_state.fecha_ref.strftime('%d/%m/%Y')}</p>", unsafe_allow_html=True)
        
        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("Vencidos o Vencen Hoy", int(vencidos))
        with col2:
            st.metric("Próximos a Vencer (1-2 días)", int(riesgo))
        with col3:
            st.metric("En Plazo (>= 3 días)", int(a_tiempo))
            
        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown("<p style='font-size: 16px; font-weight: 600; color: #12402A;'>Detalle de Registros y Control de Plazos</p>", unsafe_allow_html=True)
        
        def apply_executive_colors(row):
            try:
                dias = int(row['Días Restantes'])
                if dias <= 0:
                    return ['background-color: #FEE2E2; color: #991B1B; font-weight: 500;'] * len(row)
                elif 1 <= dias <= 2:
                    return ['background-color: #FEF3C7; color: #92400E; font-weight: 500;'] * len(row)
                elif dias >= 3:
                    return ['background-color: #ECFDF5; color: #065F46;'] * len(row)
            except:
                pass
            return [''] * len(row)

        styled_df = df_res.style.apply(apply_executive_colors, axis=1).format({'Días Restantes': '{:d}'})
        st.dataframe(styled_df, use_container_width=True, height=450, hide_index=True)
        
        st.markdown("<br>", unsafe_allow_html=True)
        col_exp1, col_exp2 = st.columns([3, 1])
        with col_exp2:
            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                styled_df.to_excel(writer, index=False, sheet_name='Control_Bloqueos')
            excel_data = output.getvalue()

            st.download_button(
                label="Descargar Reporte Excel",
                data=excel_data,
                file_name=f"Control_Bloqueos_{st.session_state.fecha_ref}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
    else:
        st.info("Cargue un archivo en la barra lateral para iniciar el procesamiento de control aduanero.")


# ==========================================
# ENRUTADOR PRINCIPAL
# ==========================================
if opcion_modulo == "Procesador de DIM":
    modulo_procesador_dim()
elif opcion_modulo == "Actas de Tránsito / PICIZ":
    modulo_actas_transito()
elif opcion_modulo == "Control Bloqueo Formularios":
    modulo_control_bloqueos()
